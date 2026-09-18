import io
import logging

import openpyxl

from app.companies.imports import get_facets, import_file
from app.models.company import Company, CompanyCandidate


def _wb_bytes(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Название", "Категории", "Регион", "Город", "Сайт", "Оценок", "Отзывов", "Рейтинг"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _wb_bytes_with_phone(rows: list[list]) -> bytes:
    """Как _wb_bytes, но с колонкой «Немобильные» — источник поля phone
    (см. app/companies/import_xlsx.py COLUMNS)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Название", "Категории", "Регион", "Город", "Сайт", "Немобильные",
              "Оценок", "Отзывов", "Рейтинг"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_workbook(rows: list[list]) -> bytes:
    """Полный заголовок выгрузки Яндекс.Карт — тот же, что и в
    test_companies_import_xlsx.py. Заголовок не должен расходиться между
    тестовыми файлами, иначе они проверяют разные структуры файла."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Запрос", "Название", "Категории", "Регион", "Город", "Полный адрес",
              "Мобильные", "Немобильные", "Сайт", "Email с сайта компании", "График",
              "Широта", "Долгота", "Оценок", "Отзывов", "Рейтинг",
              "Логотип", "Все телефоны"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_creates_candidates(db_session):
    data = _wb_bytes([["ООО Дом", "Дома", "Самара", "Самара", "https://dom.ru", 5, 3, 4.5]])
    imp = import_file(db_session, data, "builders.xlsx", uploaded_by_id=None)
    assert imp.row_count == 1
    assert imp.matched_count == 1
    candidates = db_session.query(CompanyCandidate).all()
    assert len(candidates) == 1
    assert candidates[0].site_key == "dom.ru"


def test_reimport_upserts_existing_candidate(db_session):
    first = _wb_bytes([["ООО Дом", "Дома", "Самара", "Самара", "https://dom.ru", 5, 3, 4.5]])
    import_file(db_session, first, "builders.xlsx", uploaded_by_id=None)

    second = _wb_bytes([["ООО Дом", "Дома", "Самара", "Самара", "https://dom.ru", 20, 15, 4.9]])
    import_file(db_session, second, "builders2.xlsx", uploaded_by_id=None)

    candidates = db_session.query(CompanyCandidate).all()
    assert len(candidates) == 1
    assert candidates[0].reviews_count == 15
    assert candidates[0].rating == 4.9


def test_facets_lists_distinct_region_and_category(db_session):
    data = _wb_bytes([
        ["ООО Дом", "Дома", "Самара", "Самара", "https://dom1.ru", 1, 1, 4.0],
        ["ООО Дом2", "Бани", "Москва", "Москва", "https://dom2.ru", 1, 1, 4.0],
    ])
    import_file(db_session, data, "builders.xlsx", uploaded_by_id=None)
    facets = get_facets(db_session, site_id=1)
    assert set(facets.regions) == {"Самара", "Москва"}
    assert set(facets.categories) == {"Дома", "Бани"}


def test_import_file_records_failed_status_on_broken_file(db_session):
    imp = import_file(db_session, b"not a real xlsx file", "broken.xlsx", uploaded_by_id=None)
    assert imp.status == "failed"
    assert imp.error_message
    assert db_session.query(CompanyCandidate).count() == 0


def test_facets_excludes_pairs_fully_taken_for_site(db_session):
    data = _wb_bytes([["ООО Дом", "Дома", "Самара", "Самара", "https://dom.ru", 1, 1, 4.0]])
    import_file(db_session, data, "builders.xlsx", uploaded_by_id=None)
    candidate = db_session.query(CompanyCandidate).one()
    db_session.add(Company(site_id=1, site_key=candidate.site_key, name="ООО Дом"))
    db_session.commit()

    facets = get_facets(db_session, site_id=1)
    assert "Дома" not in facets.categories


def test_import_survives_long_garbage_phone_value(db_session):
    """Регрессия реального инцидента: реальная выгрузка Яндекс.Карт иногда
    склеивает в ячейку телефона рекламный текст источника (233 символа на
    боевом файле). company_candidates.phone был String(50) — единственное
    такое значение роняло commit целой партии из 5214 строк."""
    long_phone = ("+7 (921) 776-79-70 Строительная компания DavHome. Строим "
                  "загородные дома под ключ в Ленинградской и Московской "
                  "области. Каркасные и железобетонные дома. Фиксированная "
                  "смета, договор, гарантия. Средний бюджет строительства "
                  "от 7 млн ₽.")
    assert len(long_phone) > 200

    data = _wb_bytes_with_phone(
        [["ООО Дом", "Дома", "Самара", "Самара", "https://dom.ru", long_phone, 5, 3, 4.5]])
    imp = import_file(db_session, data, "builders.xlsx", uploaded_by_id=None)

    assert imp.status == "parsed"
    candidate = db_session.query(CompanyCandidate).one()
    assert candidate.phone == long_phone


def test_commit_failure_is_logged(db_session, monkeypatch, caplog):
    """Раньше сбой финального db.commit() полностью проглатывался — ни одной
    строчки в логах, единственный способ узнать причину — воспроизводить
    вручную в python-шелле контейнера (как пришлось сделать при разборе
    реального инцидента). Форсируем сбой commit() напрямую, чтобы проверить
    логирование безотносительно конкретной причины (переполнение колонки,
    обрыв соединения с БД и т.п. — все они идут через один except-блок)."""
    data = _wb_bytes([["ООО Дом", "Дома", "Самара", "Самара", "https://dom.ru", 5, 3, 4.5]])

    original_commit = db_session.commit
    calls = {"n": 0}

    def flaky_commit():
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("симулированный сбой commit (например, обрыв связи с БД)")
        return original_commit()

    monkeypatch.setattr(db_session, "commit", flaky_commit)

    with caplog.at_level(logging.ERROR):
        imp = import_file(db_session, data, "builders.xlsx", uploaded_by_id=None)

    assert imp.status == "failed"
    assert imp.error_message == "не удалось сохранить компании — проверьте данные файла"
    assert "import_file" in caplog.text
    assert "builders.xlsx" in caplog.text


def test_import_stores_working_hours_and_logo(db_session):
    data = _make_workbook([
        ["застройщик", "ООО Дом", "Стройка", "Самарская область", "Самара",
         "ул. Ленина 1", "", "+7 846 000-00-00", "https://dom.ru", "i@dom.ru",
         "ежедневно, 09:00–18:00", 53.2, 50.1, 10, 5, 4.8,
         "https://avatars.mds.yandex.net/get-altay/1/XXXL", ""],
    ])
    import_file(db_session, data, "yandex.xlsx", uploaded_by_id=None)
    candidate = db_session.query(CompanyCandidate).one()
    assert candidate.working_hours == "ежедневно, 09:00–18:00"
    assert candidate.logo_url == "https://avatars.mds.yandex.net/get-altay/1/XXXL"


def test_reimport_updates_working_hours_and_logo(db_session):
    """Повторная загрузка того же файла — это и есть способ добрать новые
    колонки у уже импортированных кандидатов (см. §7 спеки)."""
    old = _make_workbook([
        ["застройщик", "ООО Дом", "Стройка", "Самарская область", "Самара",
         "ул. Ленина 1", "", "+7 846 000-00-00", "https://dom.ru", "i@dom.ru",
         "", 53.2, 50.1, 10, 5, 4.8, "", ""],
    ])
    import_file(db_session, old, "old.xlsx", uploaded_by_id=None)
    new = _make_workbook([
        ["застройщик", "ООО Дом", "Стройка", "Самарская область", "Самара",
         "ул. Ленина 1", "", "+7 846 000-00-00", "https://dom.ru", "i@dom.ru",
         "пн-пт 10:00–19:00", 53.2, 50.1, 10, 5, 4.8, "https://logo/1", ""],
    ])
    import_file(db_session, new, "new.xlsx", uploaded_by_id=None)
    candidate = db_session.query(CompanyCandidate).one()
    assert candidate.working_hours == "пн-пт 10:00–19:00"
    assert candidate.logo_url == "https://logo/1"


def test_value_longer_than_column_fails_with_company_and_field(db_session):
    """Тесты идут на SQLite, который не проверяет длину VARCHAR, а на проде
    Postgres откатывал весь файл с безликим «проверьте данные файла».
    Проверяем длины до записи: сообщение называет компанию и поле, в пул
    ничего не попадает."""
    long_city = "Самара" * 40
    assert len(long_city) > 200
    data = _wb_bytes([
        ["ООО Дом", "Дома", "Самара", "Самара", "https://dom.ru", 5, 3, 4.5],
        ["ООО Баня", "Бани", "Самара", long_city, "https://banya.ru", 5, 3, 4.5],
    ])
    imp = import_file(db_session, data, "builders.xlsx", uploaded_by_id=None)

    assert imp.status == "failed"
    assert "ООО Баня" in imp.error_message
    assert "Город" in imp.error_message
    assert "240" in imp.error_message and "200" in imp.error_message
    assert db_session.query(CompanyCandidate).count() == 0
